"""تست‌های سرویس خروجی اکسل تراکنش‌ها (:mod:`hesabyar.services.export`)."""
from __future__ import annotations

import openpyxl

from hesabyar.core import jalali
from hesabyar.db.models import Kind
from hesabyar.services import transactions
from hesabyar.services.export import export_transactions_xlsx

#: نمونه تراکنش‌ها: (نوع، مبلغ تومان، دسته، شرح)
_SAMPLES = [
    (Kind.INCOME, 500_000, "فروش", "فروش نقدی"),
    (Kind.EXPENSE, 120_000, "خرید", "خرید مواد اولیه"),
    (Kind.INCOME, 300_000, "فروش", "فروش کارتی"),
]


async def test_export_transactions_xlsx(store, tmp_path):
    """فایل اکسل ساخته می‌شود و محتوای آن با تراکنش‌ها همخوان است."""
    user_id = 12345
    await transactions.get_or_create_user(store, user_id, business_name="کافه من")

    # چند تراکنش درآمد/هزینه با زمان مشخص (aware تهران، از jalali.now())
    now = jalali.now()
    for kind, amount, category, description in _SAMPLES:
        await transactions.add_transaction(
            store,
            user_id,
            kind=kind,
            amount=amount,
            category=category,
            description=description,
            occurred_at=now,
        )

    # بازه‌ای که همه‌ی تراکنش‌ها را در بر می‌گیرد (ماه جاری تا اکنون)
    start, end = jalali.month_bounds(now)
    out_path = str(tmp_path / "report.xlsx")

    result = export_transactions_xlsx(store, user_id, start, end, out_path)

    # مسیر برگشتی همان مسیر خروجی است و فایل روی دیسک ساخته شده
    assert result == out_path
    assert (tmp_path / "report.xlsx").exists()

    # فایل با openpyxl دوباره باز می‌شود
    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    assert "تراکنش" in ws.title

    # همه‌ی مقادیر سلول‌ها را یک‌جا جمع می‌کنیم تا ساده بررسی کنیم
    all_cells = [cell for row_cells in ws.iter_rows() for cell in row_cells]
    all_values = [cell.value for cell in all_cells]

    # سطر هدر شامل «مبلغ (تومان)» است
    assert "مبلغ (تومان)" in all_values

    # تعداد ردیف‌های داده (ستون «نوع») برابر تعداد تراکنش‌هاست
    kind_labels = {"درآمد", "هزینه"}
    type_column_values = [row_cells[2].value for row_cells in ws.iter_rows()]
    data_row_count = sum(1 for value in type_column_values if value in kind_labels)
    assert data_row_count == len(_SAMPLES)

    # دست‌کم یکی از مبالغ به‌صورت عددی (نه رشته) در شیت هست
    numeric_values = [
        cell.value for cell in all_cells if isinstance(cell.value, int)
    ]
    assert 500_000 in numeric_values
