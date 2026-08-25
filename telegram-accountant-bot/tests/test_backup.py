import os

from openpyxl import load_workbook

from hesabyar.core import jalali
from hesabyar.db.models import Direction, Kind
from hesabyar.services import backup
from hesabyar.services import invoices as inv
from hesabyar.services import ledger
from hesabyar.services import transactions as tx

UID = 55


class TestFullExport:
    async def test_multi_sheet_export(self, store, tmp_path):
        user = await tx.get_or_create_user(store, UID)
        user.business_name = "کسب‌وکار من"
        await tx.add_transaction(
            store, UID, kind=Kind.INCOME, amount=1_000_000,
            category="فروش کالا", description="فروش", occurred_at=jalali.now(),
        )
        await tx.add_transaction(
            store, UID, kind=Kind.EXPENSE, amount=300_000,
            category="اجاره", description="اجاره مغازه", occurred_at=jalali.now(),
        )
        await ledger.add_entry(
            store, UID, direction=Direction.RECEIVABLE,
            party_name="علی", amount=500_000,
        )
        await inv.create_invoice(
            store, UID, customer_name="رضا",
            items=[{"title": "کالا", "quantity": 1, "unit_price": 200_000}],
            issue_date=jalali.now().date(),
        )

        out = str(tmp_path / "full.xlsx")
        backup.export_full_user_xlsx(store, UID, out, business=user)
        assert os.path.exists(out)

        wb = load_workbook(out)
        assert set(wb.sheetnames) == {"تراکنش‌ها", "طلب و بدهی", "فاکتورها"}
        # ۲ تراکنش (به‌علاوه‌ی سطر هدر)
        assert wb["تراکنش‌ها"].max_row == 3
        assert wb["طلب و بدهی"].max_row == 2   # هدر + ۱ ردیف
        assert wb["فاکتورها"].max_row == 2     # هدر + ۱ فاکتور

    async def test_empty_user(self, store, tmp_path):
        await tx.get_or_create_user(store, UID)
        out = str(tmp_path / "empty.xlsx")
        backup.export_full_user_xlsx(store, UID, out)
        wb = load_workbook(out)
        assert wb["تراکنش‌ها"].max_row == 1  # فقط هدر
