"""سرویس خروجی اکسل (.xlsx) از تراکنش‌ها.

این ماژول با کمک :mod:`openpyxl` یک فایل اکسل راست‌به‌چپ می‌سازد که فهرست
تراکنش‌های درآمد و هزینه‌ی یک کاربر را در یک بازه‌ی زمانی نمایش می‌دهد و در
پایان یک بخش جمع‌بندی (جمع درآمد، جمع هزینه و مانده) دارد.

همه‌ی مبالغ به «تومان» و به‌صورت عدد صحیح نوشته می‌شوند تا در اکسل عددی بمانند
و بتوان روی‌شان محاسبه انجام داد.
"""
from __future__ import annotations

import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Font
from ..db.store import Store

from ..core import jalali
from ..db.models import Kind, User
from . import transactions

#: سطر هدر جدول تراکنش‌ها
_HEADERS = ["ردیف", "تاریخ", "نوع", "دسته", "شرح", "مبلغ (تومان)"]

#: عرض معقول هر ستون (بر اساس حرف ستون)
_COLUMN_WIDTHS = {"A": 8, "B": 14, "C": 10, "D": 16, "E": 34, "F": 16}


def export_transactions_xlsx(
    store: Store,
    user_id: int,
    start: dt.datetime,
    end: dt.datetime,
    out_path: str,
    business: User | None = None,
) -> str:
    """خروجی اکسل تراکنش‌های کاربر در بازه‌ی ``[start, end]`` را می‌سازد.

    یک فایل ``.xlsx`` راست‌به‌چپ با یک شیت به نام «تراکنش‌ها» می‌سازد که شامل
    سطر هدر، یک ردیف برای هر تراکنش و یک بخش جمع‌بندی در پایان است. مسیر فایل
    ذخیره‌شده (همان ``out_path``) را برمی‌گرداند.

    پارامترها:
        store: لایه‌ی داده (Store).
        user_id: شناسه‌ی عددی کاربر (تلگرام).
        start: زمان آغاز بازه (aware، منطقه‌ی تهران).
        end: زمان پایان بازه (aware، منطقه‌ی تهران).
        out_path: مسیر ذخیره‌ی فایل اکسل.
        business: کاربر/کسب‌وکار برای درج عنوان بالای جدول (اختیاری).
    """
    txs = transactions.list_transactions(store, user_id, start, end)

    wb = Workbook()
    ws = wb.active
    ws.title = "تراکنش‌ها"
    ws.sheet_view.rightToLeft = True  # نمایش راست‌به‌چپ برای متن فارسی

    bold = Font(bold=True)
    row = 1

    # سطر اختیاری عنوان کسب‌وکار (ساده، در ستون نخست)
    if business is not None and business.business_name:
        title_cell = ws.cell(row=row, column=1, value=business.business_name)
        title_cell.font = Font(bold=True, size=14)
        row += 2  # پس از عنوان یک سطر خالی فاصله می‌گذاریم

    # سطر هدر با فونت پررنگ
    for col, title in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.font = bold
    row += 1

    # یک ردیف برای هر تراکنش
    for index, tx in enumerate(txs, start=1):
        kind_label = "درآمد" if tx.kind == Kind.INCOME else "هزینه"
        ws.cell(row=row, column=1, value=index)
        ws.cell(row=row, column=2, value=jalali.format_date(tx.occurred_at))
        ws.cell(row=row, column=3, value=kind_label)
        ws.cell(row=row, column=4, value=tx.category)
        ws.cell(row=row, column=5, value=tx.description)
        # مبلغ را عدد صحیح می‌گذاریم تا در اکسل عددی بماند (نه رشته)
        ws.cell(row=row, column=6, value=int(tx.amount))
        row += 1

    # بخش جمع‌بندی: جمع درآمد، جمع هزینه و مانده در سه سطر جداگانه
    stats = transactions.summary(store, user_id, start, end)
    row += 1  # یک سطر خالی فاصله پیش از جمع‌بندی
    for label, value in (
        ("جمع درآمد", stats["income"]),
        ("جمع هزینه", stats["expense"]),
        ("مانده", stats["balance"]),
    ):
        label_cell = ws.cell(row=row, column=5, value=label)
        label_cell.font = bold
        amount_cell = ws.cell(row=row, column=6, value=int(value))
        amount_cell.font = bold
        row += 1

    # تنظیم عرض معقول ستون‌ها
    for col_letter, width in _COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(out_path)
    return out_path
