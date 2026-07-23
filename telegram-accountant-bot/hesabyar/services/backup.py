"""پشتیبان‌گیری داده‌ها.

* :func:`export_full_user_xlsx` — اکسل چندشیتی از داده‌های یک کاربر (از Store).
* :func:`download_spreadsheet_xlsx` — دانلود کل اسپردشیت گوگل به‌صورت xlsx
  با Drive API (پشتیبان کاملِ عملیاتی برای مدیران).
"""
from __future__ import annotations

import json
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font

from ..core import jalali
from ..db.models import Direction, Kind
from ..db.store import Store

_HEADER_FONT = Font(bold=True)

_DRIVE_SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _write_header(ws, headers: list[str]) -> None:
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = _HEADER_FONT


def export_full_user_xlsx(
    store: Store, user_id: int, out_path: str, business=None
) -> str:
    """کل داده‌های کاربر را در یک فایل اکسل چندشیتی می‌نویسد."""
    wb = Workbook()

    ws = wb.active
    ws.title = "تراکنش‌ها"
    _write_header(ws, ["ردیف", "تاریخ", "نوع", "دسته", "شرح", "مبلغ (تومان)"])
    txs = sorted(
        store.list("transactions", lambda t: t.user_id == user_id),
        key=lambda t: (t.occurred_at or t.created_at, t.id),
    )
    for index, tx in enumerate(txs, start=1):
        kind_label = "درآمد" if tx.kind == Kind.INCOME else "هزینه"
        ws.append([
            index, jalali.format_date(tx.occurred_at), kind_label,
            tx.category, tx.description, int(tx.amount),
        ])

    ws2 = wb.create_sheet("طلب و بدهی")
    _write_header(ws2, ["ردیف", "نوع", "طرف‌حساب", "مبلغ (تومان)", "سررسید", "وضعیت"])
    entries = sorted(
        store.list("ledger_entries", lambda e: e.user_id == user_id),
        key=lambda e: e.id,
    )
    for index, entry in enumerate(entries, start=1):
        direction = "طلب" if entry.direction == Direction.RECEIVABLE else "بدهی"
        due = jalali.format_date(entry.due_date) if entry.due_date else "—"
        state = "تسویه‌شده" if entry.is_settled else "باز"
        ws2.append([index, direction, entry.party_name, int(entry.amount), due, state])

    ws3 = wb.create_sheet("فاکتورها")
    _write_header(ws3, ["شماره", "تاریخ", "مشتری", "جمع کل (تومان)"])
    invoices = sorted(
        store.list("invoices", lambda i: i.user_id == user_id), key=lambda i: i.seq
    )
    for inv in invoices:
        inv.items = sorted(
            store.list("invoice_items", lambda it: it.invoice_id == inv.id),
            key=lambda it: it.id,
        )
        ws3.append([
            inv.number, jalali.format_date(inv.issue_date),
            inv.customer_name, int(inv.total),
        ])

    for sheet in (ws, ws2, ws3):
        for column_cells in sheet.columns:
            length = max((len(str(c.value)) for c in column_cells if c.value), default=8)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(
                max(length + 2, 10), 40
            )

    wb.save(out_path)
    return out_path


def download_spreadsheet_xlsx(settings, out_path: str) -> Optional[str]:
    """کل اسپردشیت گوگل را به‌صورت xlsx دانلود می‌کند (Drive API).

    اگر تنظیمات گوگل ناقص باشد یا خطایی رخ دهد، ``None`` برمی‌گرداند.
    """
    if not settings.google_service_account_json or not settings.google_sheet_id:
        return None
    try:
        from google.oauth2.service_account import Credentials
        from googleapiclient.discovery import build

        info = json.loads(settings.google_service_account_json)
        creds = Credentials.from_service_account_info(info, scopes=_DRIVE_SCOPES)
        service = build("drive", "v3", credentials=creds, cache_discovery=False)
        data = service.files().export(
            fileId=settings.google_sheet_id, mimeType=_XLSX_MIME
        ).execute()
        with open(out_path, "wb") as fh:
            fh.write(data)
        return out_path
    except Exception:  # noqa: BLE001
        return None
